from pathlib import Path

import pytest
from docx import Document

from app.parsers import get_parser, parse
from app.parsers.base import Heading, Paragraph, ParserError
from app.parsers.docx_parser import DocxParser


def _build_docx(tmp_path, name: str = "doc.docx") -> Path:
    doc = Document()
    doc.add_heading("Report Title", 0)
    doc.add_heading("Section One", 1)
    doc.add_paragraph("First body paragraph.")
    doc.add_heading("Sub Section", 2)
    doc.add_paragraph("Second body paragraph with details.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Name"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "alpha"
    table.cell(1, 1).text = "42"
    path = tmp_path / name
    doc.save(str(path))
    return path


def test_registry_resolves_docx(tmp_path):
    _build_docx(tmp_path)
    assert isinstance(get_parser("docx"), DocxParser)
    assert isinstance(get_parser(".docx"), DocxParser)


def test_heading_styles_become_blocks(tmp_path):
    parsed = parse("docx", _build_docx(tmp_path))

    headings = [b for b in parsed.blocks if isinstance(b, Heading)]
    assert headings == [
        Heading(level=1, text="Report Title"),
        Heading(level=1, text="Section One"),
        Heading(level=2, text="Sub Section"),
    ]


def test_body_paragraphs_in_order(tmp_path):
    parsed = parse("docx", _build_docx(tmp_path))

    paragraphs = [b for b in parsed.blocks if isinstance(b, Paragraph)]
    assert paragraphs[0].text == "First body paragraph."
    assert paragraphs[1].text == "Second body paragraph with details."


def test_tables_flattened_to_text(tmp_path):
    parsed = parse("docx", _build_docx(tmp_path))

    texts = [b.text for b in parsed.blocks]
    assert "Name | Value" in texts
    assert "alpha | 42" in texts


def test_blocks_keep_document_order(tmp_path):
    parsed = parse("docx", _build_docx(tmp_path))

    assert isinstance(parsed.blocks[0], Heading)
    assert parsed.blocks[0].text == "Report Title"
    assert any(b.text == "First body paragraph." for b in parsed.blocks)


def test_full_text_concatenates_blocks(tmp_path):
    parsed = parse("docx", _build_docx(tmp_path))

    assert "Report Title" in parsed.text
    assert "First body paragraph." in parsed.text
    assert "Name | Value" in parsed.text


def test_corrupt_file_raises_parser_error(tmp_path):
    path = tmp_path / "bad.docx"
    path.write_bytes(b"this is not a docx file at all")

    with pytest.raises(ParserError):
        parse("docx", path)
